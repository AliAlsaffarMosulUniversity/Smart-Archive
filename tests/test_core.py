import tempfile, unittest
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from core import *
class ArchiveTests(unittest.TestCase):
    def setUp(self):
        self.t=tempfile.TemporaryDirectory(); self.base=Path(self.t.name)
        self.src=self.base/'original.txt'; self.src.write_text('الجهة المصدرة: جامعة الموصل\nإلى: جامعة بغداد\nالتاريخ: 2025/03/12\nالموضوع: شكر وتقدير\nالعدد: 123\nموعد الدورة 2026/11/15',encoding='utf8')
        self.p=propose(self.src,*extract(self.src)); self.a=Archive(self.base/'archive')
    def tearDown(self): self.a.close(); self.t.cleanup()
    def test_fields_and_issue_date(self):
        self.assertEqual(self.p['category'],'شكر وتقدير'); self.assertEqual(self.p['date'],'2025-03-12'); self.assertEqual(self.p['issuer'],'جامعة الموصل')
    def test_copy_dedupe_undo(self):
        dest,dup=self.a.store(self.p); self.assertFalse(dup); self.assertEqual(digest(dest),digest(self.src)); self.assertTrue(self.src.exists())
        self.assertTrue(self.a.store(self.p)[1]); ident=self.a.search()[0][0]; self.a.undo(ident)
        self.assertFalse(Path(dest).exists()); self.assertTrue(self.src.exists()); self.assertEqual(self.a.search()[0][-1],'undone')
    def test_changed_source_rejected(self):
        self.src.write_text('changed'); self.assertRaises(ValueError,self.a.store,self.p)
    def test_changed_archive_not_removed(self):
        dest,_=self.a.store(self.p); Path(dest).write_text('edited'); self.assertRaises(ValueError,self.a.undo,self.a.search()[0][0]); self.assertTrue(Path(dest).exists())
    def test_missing_date_not_invented(self):
        p=propose(self.src,[{'location':'1','text':'دورة تدريبية بتاريخ 2026/10/25'}]); self.assertEqual(p['year'],'غير محدد')
    def test_invalid_date(self):
        p=propose(self.src,[{'location':'1','text':'التاريخ: 2026/02/31'}]); self.assertEqual(p['date'],'')
    def test_sanitize_and_traversal(self):
        self.p['issuer']='../../CON'; dest,_=self.a.store(self.p); self.assertTrue(Path(dest).resolve().is_relative_to(self.a.root)); self.assertEqual(safe('CON'),'_CON')
    def test_arabic_search(self):
        self.a.store(self.p); self.assertTrue(self.a.search('شُكر')); self.assertTrue(self.a.search('الموصل'))
    def test_unsupported(self):
        bad=self.base/'x.doc'; bad.write_bytes(b'x'); self.assertRaises(ValueError,extract,bad)
    def test_excel_csv_injection(self):
        self.p['subject']='=HYPERLINK("bad")'; self.a.store(self.p); out=self.base/'export.csv'; self.a.export(out); self.assertIn("'=HYPERLINK",out.read_text(encoding='utf-8-sig'))
    def test_ooxml(self):
        from docx import Document
        d=Document(); d.add_paragraph('الموضوع: شكر وتقدير'); path=self.base/'x.docx'; d.save(path); self.assertIn('شكر',extract(path)[0][0]['text'])
        from openpyxl import Workbook
        w=Workbook(); w.active.append(['التعليم المستمر']); path=self.base/'x.xlsx'; w.save(path); self.assertIn('المستمر',extract(path)[0][0]['text'])
        from pptx import Presentation
        r=Presentation(); s=r.slides.add_slide(r.slide_layouts[1]); s.shapes.title.text='تشكيل لجنة'; path=self.base/'x.pptx'; r.save(path); self.assertIn('لجنة',extract(path)[0][0]['text'])
if __name__=='__main__': unittest.main()
